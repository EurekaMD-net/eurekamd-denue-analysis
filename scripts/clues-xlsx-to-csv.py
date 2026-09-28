#!/usr/bin/env python3
r"""
Convert the DGIS CLUES catalogue XLSX to the CSV that load-clues.ts consumes.

Source: http://gobi.salud.gob.mx/gobi/catalogos/catalogosmaestros/
        ESTABLECIMIENTO_SALUD_YYYYMM.xlsx (sheets CLUES_YYYYMM, SUBCLUES_YYYYMM,
        HORARIOS_YYYYMM). Only the first CLUES_ sheet is converted (68 columns).

Why Python: \copy can't read XLSX (binary). Same openpyxl read_only pattern as
coneval-ageb-xlsx-to-csv.py; replaces the ad-hoc snippet used for the 2026-04
cut, which was never checked in.

Contract (fail closed, exit 1): the snake_cased source headers must equal
HEADER (= CLUES_STAGING_DDL in load-clues.ts; the loader's \copy is
positional); every cell must be text (DGIS ships codes as zero-padded strings
and dates as dd/mm/yyyy text; a numeric cell would silently break the cve_mun
join); no row wider than 68 (short rows are padded); at least MIN_ROWS rows.
Empty fields load as NULL under `\copy ... WITH (FORMAT csv, HEADER true)`.

Usage:
  python3 scripts/clues-xlsx-to-csv.py \
      --xlsx=raw/clues/ESTABLECIMIENTO_SALUD_202608.xlsx \
      --out=raw/clues/clues_202608.csv
"""

import csv
import os
import re
import sys

import openpyxl


# Mirror of CLUES_STAGING_DDL in scripts/load-clues.ts — keep in sync.
HEADER = [
    "clues",
    "clave_de_la_institucion", "nombre_de_la_institucion",
    "clave_de_la_entidad", "entidad",
    "clave_del_municipio", "municipio",
    "clave_de_la_localidad", "localidad",
    "clave_de_la_jurisdiccion", "jurisdiccion",
    "clave_del_tipo_establecimiento", "nombre_tipo_establecimiento",
    "clave_de_tipologia", "nombre_de_tipologia",
    "clave_de_subtipologia", "nombre_de_subtipologia",
    "nombre_de_la_unidad", "nombre_comercial",
    "clave_tipo_de_vialidad", "tipo_de_vialidad", "vialidad",
    "numero_exterior", "numero_interior",
    "clave_tipo_de_asentamiento", "tipo_de_asentamiento", "asentamiento",
    "codigo_postal", "referencias_del_domicilio",
    "clave_estatus_de_operacion", "estatus_de_operacion",
    "rfc_del_establecimiento",
    "telefono_1_del_establecimiento", "extension_telefonica_1_del_establecimiento",
    "telefono_2_del_establecimiento", "extension_telefonica_2_del_establecimiento",
    "fecha_de_construccion", "fecha_de_inicio_de_operacion",
    "clave_unidad_movil_marca", "unidad_movil_marca",
    "unidad_movil_marca_especifica", "unidad_movil_modelo",
    "clave_unidad_movil_programa", "unidad_movil_programa",
    "clave_unidad_movil_tipo", "unidad_movil_tipo",
    "clave_unidad_movil_tipologia", "unidad_movil_tipologia",
    "clave_de_la_ins_adm", "nombre_de_la_ins_adm",
    "clave_nivel_atencion", "nivel_atencion",
    "clave_estrato_unidad", "estrato_unidad",
    "clave_tipo_obra", "tipo_obra",
    "clave_propiedad_del_inmueble", "propiedad_del_inmueble",
    "observaciones_al_registro",
    "latitud", "longitud",
    "clave_ultimo_movimiento", "ultimo_movimiento", "fecha_ultimo_movimiento",
    "comentarios_de_la_validacion",
    "clave_motivo_baja", "motivo_baja", "fecha_efectiva_de_baja",
]
MIN_ROWS = 60_000  # source cuts carry ~63-65k rows
USAGE = "usage: clues-xlsx-to-csv.py --xlsx=<path> --out=<path>\n"


def fail(msg, code=1):
    sys.stderr.write(f"clues-xlsx-to-csv: {msg}\n")
    sys.exit(code)


def main():
    args = dict(a.partition("=")[::2] for a in sys.argv[1:])
    if len(sys.argv) != 3 or set(args) != {"--xlsx", "--out"} or not all(args.values()):
        fail(f"bad arguments {sys.argv[1:]}\n{USAGE}", 2)

    wb = openpyxl.load_workbook(args["--xlsx"], read_only=True, data_only=True)
    sheet = next((s for s in wb.sheetnames if s.startswith("CLUES_")), None)
    if sheet is None:
        fail(f"no CLUES_ sheet; workbook has {wb.sheetnames}")
    ws = wb[sheet]
    ws.reset_dimensions()  # a stale <dimension> tag silently truncates reads
    rows = ws.iter_rows(values_only=True)

    header = [re.sub(r"\s+", "_", str(h or "").strip()).lower() for h in next(rows, ())]
    if header != HEADER:
        i = next((i for i, (a, b) in enumerate(zip(header, HEADER)) if a != b), min(len(header), 68))
        fail(f"header mismatch at index {i}: got {header[i:i + 1]}, expected {HEADER[i:i + 1]}")

    tmp = args["--out"] + ".tmp"
    n = 0
    try:
        with open(tmp, "w", encoding="utf-8", newline="") as f:
            out = csv.writer(f, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
            out.writerow(HEADER)
            for rownum, row in enumerate(rows, start=2):
                if len(row) > len(HEADER):
                    fail(f"row {rownum} has {len(row)} cells (> {len(HEADER)})")
                for j, v in enumerate(row):
                    if v is not None and not isinstance(v, str):
                        fail(f"row {rownum} column {HEADER[j]}: non-text cell {v!r}")
                out.writerow([v or "" for v in row] + [""] * (len(HEADER) - len(row)))
                n += 1
        if n < MIN_ROWS:
            fail(f"only {n} data rows (< {MIN_ROWS}); truncated sheet?")
        os.replace(tmp, args["--out"])
    except BaseException:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise

    print(f"clues-xlsx-to-csv: rows={n} columns={len(HEADER)} sheet={sheet}")


if __name__ == "__main__":
    main()
